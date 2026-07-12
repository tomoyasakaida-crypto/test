# -*- coding: utf-8 -*-
"""
Excel の週次更新（upsert）。openpyxl のみ使用。
 - 物件一覧シート … 主キーで更新。新規/継続/掲載終了を追跡し、履歴を残す。
 - 入居日不明シート … 入居可能日が判定できなかった物件（レビュー用）。
 - 実行ログシート … 毎回の取得日と件数。
"""
import os
from datetime import date
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

SHEET_MAIN = "物件一覧"
SHEET_UNKNOWN = "入居日不明"
SHEET_LOG = "実行ログ"

COLUMNS = [
    "主キー", "建物名", "住所", "坪数", "賃料",
    "入居可能日(記載)", "入居可能日(解析)", "入居可能日区分",
    "条件①3ヶ月以上先", "条件②大阪市", "条件③30坪以上",
    "詳細URL", "初回取得日", "最終確認日", "ステータス",
]
_KEY = 0
_FIRST_SEEN = COLUMNS.index("初回取得日")
_LAST_SEEN = COLUMNS.index("最終確認日")
_STATUS = COLUMNS.index("ステータス")

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_NEW_FILL = PatternFill("solid", fgColor="E2EFDA")
_GONE_FILL = PatternFill("solid", fgColor="F2F2F2")


def make_key(rec: dict) -> str:
    url = (rec.get("detail_url") or "").split("?")[0].rstrip("/")
    if url:
        return url
    return "|".join([rec.get("name", ""), rec.get("address", ""), str(rec.get("tsubo", ""))])


def _bool_mark(v):
    return {True: "○", False: "×", None: "△"}.get(v, "")


def record_to_row(rec: dict, j: dict, today: date):
    return [
        make_key(rec), rec.get("name", ""), rec.get("address", ""),
        rec.get("tsubo", ""), rec.get("rent", ""),
        rec.get("move_in_text", ""), j.get("move_in_date", ""), j.get("move_in_kind", ""),
        _bool_mark(j.get("cond1")), _bool_mark(j.get("cond2")), _bool_mark(j.get("cond3")),
        rec.get("detail_url", ""), today.isoformat(), today.isoformat(), "新規",
    ]


def _ensure_wb(path):
    if os.path.exists(path):
        return load_workbook(path)
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_MAIN
    _write_header(ws)
    wb.create_sheet(SHEET_UNKNOWN)
    _write_header(wb[SHEET_UNKNOWN])
    log = wb.create_sheet(SHEET_LOG)
    log.append(["取得日", "対象(3条件合致)", "新規", "継続", "掲載終了", "入居日不明"])
    for c in log[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = _HEADER_FILL
    return wb


def _write_header(ws):
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = _HEADER_FILL
        c.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"


def _index_existing(ws):
    idx = {}
    for r in range(2, ws.max_row + 1):
        key = ws.cell(r, 1).value
        if key:
            idx[key] = r
    return idx


def update_workbook(path, matched, unknown, today: date):
    """
    matched: 3条件合致した物件 [(rec, judge_dict), ...]
    unknown: 入居日不明で退避する物件 [(rec, judge_dict), ...]
    """
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb = _ensure_wb(path)
    counts = _upsert_sheet(wb[SHEET_MAIN], matched, today)
    _upsert_sheet(wb[SHEET_UNKNOWN], unknown, today, track_gone=False)

    log = wb[SHEET_LOG]
    log.append([today.isoformat(), len(matched), counts["new"], counts["kept"],
                counts["gone"], len(unknown)])

    _autosize(wb[SHEET_MAIN])
    _autosize(wb[SHEET_UNKNOWN])
    wb.save(path)
    return counts


def _upsert_sheet(ws, items, today: date, track_gone=True):
    existing = _index_existing(ws)
    seen_keys = set()
    new = kept = 0

    for rec, j in items:
        row = record_to_row(rec, j, today)
        key = row[_KEY]
        seen_keys.add(key)
        if key in existing:
            r = existing[key]
            first_seen = ws.cell(r, _FIRST_SEEN + 1).value or today.isoformat()
            for i, val in enumerate(row, start=1):
                ws.cell(r, i, val)
            ws.cell(r, _FIRST_SEEN + 1, first_seen)   # 初回取得日は維持
            ws.cell(r, _STATUS + 1, "継続")
            kept += 1
        else:
            ws.append(row)
            for c in ws[ws.max_row]:
                c.fill = _NEW_FILL
            new += 1

    gone = 0
    if track_gone:
        for key, r in existing.items():
            if key not in seen_keys:
                cur = ws.cell(r, _STATUS + 1).value
                if cur != "掲載終了":
                    ws.cell(r, _STATUS + 1, "掲載終了")
                    for c in ws[r]:
                        c.fill = _GONE_FILL
                gone += 1
    return {"new": new, "kept": kept, "gone": gone}


def _autosize(ws, cap=48):
    for col in ws.columns:
        width = 8
        letter = col[0].column_letter
        for c in col:
            if c.value is not None:
                width = max(width, min(cap, len(str(c.value)) + 2))
        ws.column_dimensions[letter].width = width
